"""Generate the synthetic CV fixtures used by the extraction tests.

Every name, company and contact detail here is invented. Real CVs must never be
used as fixtures (they live only in gitignored data/private/).

Run from the repo root:  uv run python tests/extraction/make_fixtures.py
Output is deterministic, so re-running it on a clean checkout leaves git clean.
"""

from __future__ import annotations

import io
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import docx
import pymupdf
from docx.oxml import parse_xml
from docx.shared import Pt

FIXTURES_DIR = Path(__file__).parent / "fixtures"

# Fixed metadata so saved bytes do not depend on when the script ran.
_METADATA = {
    "title": "Synthetic CV fixture",
    "author": "JobPDF test suite",
    "creator": "tests/extraction/make_fixtures.py",
    "producer": "JobPDF",
    "creationDate": "D:20260101000000Z",
    "modDate": "D:20260101000000Z",
}

_FIXED_TIME = datetime(2026, 1, 1, tzinfo=timezone.utc)
_ZIP_TIME = (2026, 1, 1, 0, 0, 0)

# A legacy VML floating text box. Its text must NOT reach full_text; the parser
# only warns about it.
_TEXT_BOX_RUN = (
    '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    ' xmlns:v="urn:schemas-microsoft-com:vml">'
    '<w:pict><v:shape style="width:200pt;height:40pt"><v:textbox><w:txbxContent>'
    "<w:p><w:r><w:t>Floating note text</w:t></w:r></w:p>"
    "</w:txbxContent></v:textbox></v:shape></w:pict></w:r>"
)

BODY = 10.0
HEADER = 14.0
NAME = 22.0


def _text(page: pymupdf.Page, x: float, y: float, lines: list[str], *,
          size: float = BODY, bold: bool = False, gap: float = 1.0) -> float:
    """Write lines with baseline starting at y; return the y for the next block.

    ``gap`` is extra blank space (in line heights) after the block, which is what
    makes MuPDF split consecutive groups of lines into separate blocks.
    """
    font = "hebo" if bold else "helv"
    leading = size * 1.25
    for line in lines:
        page.insert_text((x, y), line, fontname=font, fontsize=size)
        y += leading
    return y + leading * gap


def _spaced_items(page: pymupdf.Page, x: float, y: float, items: list[str]) -> float:
    """Write one item per block, like sidebar skill lists with generous spacing.

    Separate blocks are what a naive top-to-bottom sort interleaves with the
    neighbouring column, so this is what the contiguity test needs.
    """
    for item in items:
        y = _text(page, x, y, [item], gap=0.6)
    return y


def _save(doc: pymupdf.Document, name: str) -> None:
    doc.set_metadata(_METADATA)
    doc.save(FIXTURES_DIR / name, garbage=4, deflate=True, no_new_id=True)
    doc.close()


def single_column() -> pymupdf.Document:
    doc = pymupdf.open()
    page = doc.new_page()  # A4-ish default 595 x 842
    x = 50
    y = _text(page, x, 70, ["Olena Testenko"], size=NAME, bold=True, gap=0.3)
    y = _text(page, x, y, ["olena.testenko@example.com | +380 00 000 0000 | Kyiv, Ukraine"])
    y = _text(page, x, y, ["Experience"], size=HEADER, bold=True, gap=0.4)
    y = _text(page, x, y, ["Data Engineer"], bold=True, gap=0.2)
    y = _text(page, x, y, ["Acme Widgets LLC, 2021-2024"], gap=0.4)
    y = _text(page, x, y, [
        "- Led the develop-",
        "ment of nightly reporting jobs for the finance team.",
        "- Reduced report delivery time from hours to minutes.",
    ])
    y = _text(page, x, y, ["Junior Analyst"], bold=True, gap=0.2)
    y = _text(page, x, y, ["Globex Testing Co, 2019-2021"], gap=0.4)
    y = _text(page, x, y, ["- Maintained weekly sales dashboards for regional managers."])
    y = _text(page, x, y, ["Education"], size=HEADER, bold=True, gap=0.4)
    y = _text(page, x, y, ["BSc Computer Science, Example State University, 2015-2019"])
    y = _text(page, x, y, ["Skills"], size=HEADER, bold=True, gap=0.4)
    _text(page, x, y, ["Python, SQL, Apache Airflow, Docker, PostgreSQL"])
    return doc


def two_column() -> pymupdf.Document:
    doc = pymupdf.open()
    page = doc.new_page()
    # Full-width banner: short centred name plus a wide contact line.
    _text(page, 225, 60, ["Olena Testenko"], size=NAME, bold=True)
    _text(page, 50, 92, [
        "olena.testenko@example.com | +380 00 000 0000 | Kyiv, Ukraine | example.com/olena-demo",
    ])

    left_x, right_x, top = 50, 310, 140
    y = _text(page, left_x, top, ["Skills"], size=HEADER, bold=True, gap=0.4)
    y = _spaced_items(page, left_x, y, [
        "Python", "SQL", "Apache Airflow", "Docker", "PostgreSQL",
    ])
    y = _text(page, left_x, y, ["Languages"], size=HEADER, bold=True, gap=0.4)
    _text(page, left_x, y, ["Ukrainian (native)", "English (C1)"])

    y = _text(page, right_x, top, ["Experience"], size=HEADER, bold=True, gap=0.4)
    y = _text(page, right_x, y, ["Data Engineer"], bold=True, gap=0.2)
    y = _text(page, right_x, y, ["Acme Widgets LLC, 2021-2024"], gap=0.4)
    y = _text(page, right_x, y, [
        "- Built nightly reporting jobs for finance.",
        "- Cut report delivery from hours to minutes.",
    ])
    y = _text(page, right_x, y, ["Junior Analyst"], bold=True, gap=0.2)
    y = _text(page, right_x, y, ["Globex Testing Co, 2019-2021"], gap=0.4)
    y = _text(page, right_x, y, ["- Maintained weekly sales dashboards."])
    y = _text(page, right_x, y, ["Education"], size=HEADER, bold=True, gap=0.4)
    _text(page, right_x, y, ["BSc Computer Science", "Example State University, 2015-2019"])
    return doc


def left_sidebar() -> pymupdf.Document:
    doc = pymupdf.open()
    page = doc.new_page()
    side_x, main_x, top = 30, 170, 60

    # Narrow sidebar (~25% of the page), vertically aligned with the main column
    # so a naive top-to-bottom sort would interleave the two.
    y = _text(page, side_x, top, ["Contact"], size=HEADER, bold=True, gap=0.4)
    y = _text(page, side_x, y, ["olena.t@example.com", "+380 00 000 0000", "Kyiv, Ukraine"])
    y = _text(page, side_x, y, ["Skills"], size=HEADER, bold=True, gap=0.4)
    _spaced_items(page, side_x, y, ["Python", "SQL", "Docker", "Kubernetes", "Terraform"])

    # Main column paragraphs are wider than 60% of the page on purpose.
    y = _text(page, main_x, top + 10, ["Olena Testenko"], size=NAME, bold=True, gap=0.5)
    y = _text(page, main_x, y, ["Profile"], size=HEADER, bold=True, gap=0.4)
    y = _text(page, main_x, y, [
        "Platform engineer with six years of experience running internal services for product",
        "teams of various sizes, with a strong focus on reliability and clear on-call practices",
        "that keep the incident load low for everyone who is involved in the weekly support rota.",
    ])
    y = _text(page, main_x, y, ["Experience"], size=HEADER, bold=True, gap=0.4)
    y = _text(page, main_x, y, ["Platform Engineer"], bold=True, gap=0.2)
    y = _text(page, main_x, y, ["Initech Demo Ltd, 2020-2024"], gap=0.4)
    _text(page, main_x, y, [
        "- Moved the deployment pipeline to a declarative setup that is now used by four teams,",
        "  replacing the hand-written release scripts and halving the average change lead time.",
    ])
    return doc


def image_only() -> pymupdf.Document:
    """Simulate a scanned CV: a page that is only a picture of text."""
    source = single_column()
    pix = source[0].get_pixmap(dpi=100)
    rect = source[0].rect
    source.close()

    doc = pymupdf.open()
    page = doc.new_page(width=rect.width, height=rect.height)
    page.insert_image(page.rect, pixmap=pix)
    return doc


def table_cv() -> docx.document.Document:
    """A CV built from tables, with contact details in the page header."""
    doc = docx.Document()
    props = doc.core_properties
    props.author = props.last_modified_by = "JobPDF test suite"
    props.title = "Synthetic CV fixture"
    props.created = props.modified = _FIXED_TIME
    props.revision = 1

    header = doc.sections[0].header
    header.is_linked_to_previous = False
    name = header.paragraphs[0]
    name_run = name.add_run("Olena Testenko")
    name_run.bold = True
    name_run.font.size = Pt(20)
    header.add_paragraph("olena.testenko@example.com | +380 00 000 0000 | Kyiv, Ukraine")

    doc.add_heading("Experience", level=1)
    jobs = doc.add_table(rows=3, cols=2)
    rows = [
        ("2021-2024", "Data Engineer", "Acme Widgets LLC"),
        ("2019-2021", "Junior Analyst", "Globex Testing Co"),
    ]
    for row, (dates, title, company) in zip(list(jobs.rows)[:2], rows, strict=True):
        row.cells[0].text = dates
        title_paragraph = row.cells[1].paragraphs[0]
        title_paragraph.add_run(title).bold = True
        row.cells[1].add_paragraph(company)
    merged = jobs.rows[2].cells[0].merge(jobs.rows[2].cells[1])
    merged.text = "Selected project: internal reporting portal"

    doc.add_heading("Education", level=1)
    education = doc.add_table(rows=1, cols=2)
    education.rows[0].cells[0].text = "2015-2019"
    education.rows[0].cells[1].text = "BSc Computer Science, Example State University"

    doc.add_heading("Skills", level=1)
    skills = doc.add_paragraph("Python, SQL, Docker")
    skills._p.append(parse_xml(_TEXT_BOX_RUN))
    return doc


def _save_docx(doc: docx.document.Document, name: str) -> None:
    """Save with fixed zip timestamps; python-docx stamps entries with 'now'."""
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    with zipfile.ZipFile(buffer) as src, zipfile.ZipFile(FIXTURES_DIR / name, "w") as dst:
        for item in src.infolist():
            info = zipfile.ZipInfo(item.filename, date_time=_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(info, src.read(item.filename))


def main() -> None:
    FIXTURES_DIR.mkdir(exist_ok=True)
    _save(single_column(), "single_column.pdf")
    _save(two_column(), "two_column.pdf")
    _save(left_sidebar(), "left_sidebar.pdf")
    _save(image_only(), "image_only.pdf")
    _save_docx(table_cv(), "table_cv.docx")
    (FIXTURES_DIR / "corrupt.pdf").write_bytes(b"not a pdf\x00\xff\x13 garbage\n")
    print(f"Fixtures written to {FIXTURES_DIR}")


if __name__ == "__main__":
    main()
