from pathlib import Path

import docx
import pytest
from docx.oxml import parse_xml

from jobpdf.extraction import ParseError, parse
from jobpdf.extraction.parse_docx import NO_TEXT_WARNING, TEXT_BOX_WARNING

FIXTURES = Path(__file__).parent / "fixtures"
TABLE_CV = FIXTURES / "table_cv.docx"


def test_offsets_index_into_full_text() -> None:
    doc = parse(TABLE_CV)

    assert doc.source_type == "docx"
    assert doc.has_text_layer
    for block in doc.blocks:
        assert doc.full_text[block.start : block.end] == block.text
        assert block.page is None and block.bbox is None


def test_header_contact_info_comes_before_table_body_in_order() -> None:
    text = parse(TABLE_CV).full_text

    in_reading_order = [
        "Olena Testenko",
        "olena.testenko@example.com",
        "Experience",
        "2021-2024",
        "Data Engineer",
        "Acme Widgets LLC",
        "2019-2021",
        "Junior Analyst",
        "Education",
        "BSc Computer Science",
        "Skills",
    ]
    positions = [text.index(fragment) for fragment in in_reading_order]
    assert positions == sorted(positions)


def test_merged_table_cell_appears_once() -> None:
    text = parse(TABLE_CV).full_text

    assert text.count("Selected project: internal reporting portal") == 1


def test_headings_are_bold_and_larger_than_body() -> None:
    doc = parse(TABLE_CV)

    heading = next(b for b in doc.blocks if b.text == "Experience")
    body = next(b for b in doc.blocks if b.text == "Acme Widgets LLC")
    assert heading.is_bold
    assert not body.is_bold
    assert heading.font_size and body.font_size and heading.font_size > body.font_size


def test_text_box_is_warned_about_not_extracted() -> None:
    doc = parse(TABLE_CV)

    assert TEXT_BOX_WARNING in doc.warnings
    assert "Floating note text" not in doc.full_text


def test_content_control_paragraphs_are_extracted(tmp_path: Path) -> None:
    source = docx.Document()
    source.add_paragraph("Before the control")
    source.element.body.insert(
        1,
        parse_xml(
            '<w:sdt xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            "<w:sdtContent><w:p><w:r><w:t>Inside a content control</w:t></w:r></w:p>"
            "</w:sdtContent></w:sdt>"
        ),
    )
    path = tmp_path / "sdt.docx"
    source.save(str(path))

    text = parse(path).full_text
    assert text.index("Before the control") < text.index("Inside a content control")


def test_empty_document_has_no_text(tmp_path: Path) -> None:
    path = tmp_path / "empty.docx"
    docx.Document().save(str(path))

    doc = parse(path)
    assert doc.blocks == []
    assert doc.has_text_layer is False
    assert NO_TEXT_WARNING in doc.warnings


def test_corrupt_docx(tmp_path: Path) -> None:
    path = tmp_path / "broken.docx"
    path.write_bytes(b"PK\x03\x04 definitely not a complete zip archive")

    with pytest.raises(ParseError) as exc:
        parse(path)
    assert exc.value.reason == "corrupt"
